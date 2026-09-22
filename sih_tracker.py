#!/usr/bin/env python3
"""
SIH 2026 Submission Tracker
---------------------------
Scrapes https://sih.gov.in/sih2026PS every 5 minutes and updates an Excel file
with the submitted-idea count for every problem statement.

Workbook sheets
  Live     - current snapshot of all PS, change since last run, change today
  History  - one row per PS each time its count changed (timestamped log)
  Summary  - totals, last update time, top 10 most/least crowded PS

Usage
  pip install beautifulsoup4 lxml openpyxl
  python sih_tracker.py                 # run forever, every 5 minutes
  python sih_tracker.py --once          # run a single update (for cron / Task Scheduler)
  python sih_tracker.py --interval 10   # custom interval in minutes
  python sih_tracker.py --watch SIH26011 SIH26084   # highlight your shortlisted PS
  python sih_tracker.py --cache page.html --once    # parse a saved copy of the page

Note: if the Excel file is open in Excel while an update runs, Windows locks it.
The script will then save to a timestamped copy instead and try the main file
again on the next run.
"""

import argparse
import json
import re
import shutil
import subprocess
import sys
import time
import urllib.request
from datetime import datetime
from pathlib import Path

try:
    from bs4 import BeautifulSoup
    from openpyxl import Workbook, load_workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
except ImportError:
    sys.exit("Missing packages. Run:  pip install beautifulsoup4 lxml openpyxl")

URL = "https://sih.gov.in/sih2026PS"
EXCEL_FILE = Path("SIH2026_Submissions.xlsx")
STATE_FILE = Path(".sih_tracker_state.json")

UA = ("Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 "
      "(KHTML, like Gecko) Chrome/126.0.0.0 Safari/537.36")

FONT = Font(name="Arial", size=10)
HEAD_FONT = Font(name="Arial", size=10, bold=True, color="FFFFFF")
HEAD_FILL = PatternFill("solid", fgColor="1F4E78")
UP_FILL = PatternFill("solid", fgColor="FFF2CC")       # count changed this run
WATCH_FILL = PatternFill("solid", fgColor="DDEBF7")    # your shortlisted PS

MOJIBAKE = {
    "\u00e2\u20ac\u201c": "\u2013", "\u00e2\u20ac\u201d": "\u2014",
    "\u00e2\u20ac\u2122": "\u2019", "\u00e2\u20ac\u02dc": "\u2018",
    "\u00e2\u20ac\u0153": "\u201c", "\u00e2\u20ac\u009d": "\u201d",
    "\u00c2\u00b0": "\u00b0",
}


def log(msg):
    print(f"[{datetime.now():%Y-%m-%d %H:%M:%S}] {msg}", flush=True)


def clean(text):
    for bad, good in MOJIBAKE.items():
        text = text.replace(bad, good)
    return " ".join(text.split())


# --------------------------------------------------------------------------- fetch
def fetch_html():
    """The SIH site sits behind a WAF that often blocks Python's default client,
    so try curl first (ships with Windows 10+, macOS, Linux), then urllib."""
    headers = {
        "User-Agent": UA,
        "Accept": "text/html,application/xhtml+xml,application/xml;q=0.9,*/*;q=0.8",
        "Accept-Language": "en-US,en;q=0.9",
        "Referer": "https://sih.gov.in/",
    }
    if shutil.which("curl"):
        cmd = ["curl", "-sS", "-L", "--compressed", "--max-time", "90"]
        for k, v in headers.items():
            cmd += ["-H", f"{k}: {v}"]
        cmd.append(URL)
        r = subprocess.run(cmd + ["-w", "\nHTTP_STATUS:%{http_code}"],
                           capture_output=True, timeout=120)
        out = r.stdout.decode("utf-8", errors="replace")
        body, _, status = out.rpartition("\nHTTP_STATUS:")
        if r.returncode == 0 and "dataTablePS" in body:
            return body
        log(f"curl got HTTP {status or '?'} (exit {r.returncode}): {body[:150]!r}")
    req = urllib.request.Request(URL, headers=headers)
    with urllib.request.urlopen(req, timeout=90) as resp:
        return resp.read().decode("utf-8", errors="replace")


def fetch_with_retry(attempts=3):
    for i in range(attempts):
        try:
            html = fetch_html()
            if "dataTablePS" in html:
                return html
            raise RuntimeError("page loaded but the problem-statement table is missing")
        except Exception as e:  # network errors, WAF blocks, timeouts
            wait = 15 * (i + 1)
            log(f"Fetch attempt {i + 1}/{attempts} failed: {e}")
            if i < attempts - 1:
                time.sleep(wait)
    return None


# --------------------------------------------------------------------------- parse
def to_int(text):
    m = re.search(r"\d+", text.replace(",", ""))
    return int(m.group()) if m else 0


def parse(html):
    soup = BeautifulSoup(html, "lxml")
    table = soup.find("table", id="dataTablePS")
    if not table:
        raise RuntimeError("Could not find table #dataTablePS - site layout may have changed")
    body = table.find("tbody") or table
    rows = []
    for tr in body.find_all("tr", recursive=False):
        tds = tr.find_all("td", recursive=False)
        if len(tds) < 8:
            continue
        link = tds[2].find("a")
        title = link.get_text(strip=True) if link else tds[2].get_text(strip=True)
        rows.append({
            "sno": to_int(tds[0].get_text()),
            "org": clean(tds[1].get_text(" ", strip=True)),
            "title": clean(title),
            "category": clean(tds[3].get_text(" ", strip=True)),
            "ps_id": clean(tds[4].get_text(strip=True)),
            "ideas": to_int(tds[5].get_text()),
            "theme": clean(tds[6].get_text(" ", strip=True)),
            "deadline": clean(tds[7].get_text(" ", strip=True)),
        })
    if not rows:
        raise RuntimeError("Table found but no rows parsed")
    return rows


# --------------------------------------------------------------------------- state
def load_state():
    if STATE_FILE.exists():
        try:
            return json.loads(STATE_FILE.read_text(encoding="utf-8"))
        except json.JSONDecodeError:
            pass
    return {"last": {}, "day": "", "day_start": {}, "pending_history": []}


def save_state(state):
    STATE_FILE.write_text(json.dumps(state, indent=1), encoding="utf-8")


# --------------------------------------------------------------------------- excel
def style_header(ws, headers, widths):
    ws.append(headers)
    for i, (cell, w) in enumerate(zip(ws[1], widths), start=1):
        cell.font, cell.fill = HEAD_FONT, HEAD_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A2"


def write_live(wb, rows, prev, day_start, watch, stamp):
    if "Live" in wb.sheetnames:
        del wb["Live"]
    ws = wb.create_sheet("Live", 0)
    headers = ["S.No", "PS ID", "Title", "Organisation", "Category", "Theme",
               "Submitted Ideas", "Change (last run)", "Change (today)", "Deadline", "Last Checked"]
    style_header(ws, headers, [6, 11, 60, 30, 11, 26, 11, 11, 11, 16, 18])
    for r in rows:
        before = prev.get(r["ps_id"])
        delta = r["ideas"] - before if before is not None else 0
        today = r["ideas"] - day_start.get(r["ps_id"], r["ideas"])
        ws.append([r["sno"], r["ps_id"], r["title"], r["org"], r["category"], r["theme"],
                   r["ideas"], delta, today, r["deadline"], stamp])
        fill = WATCH_FILL if r["ps_id"] in watch else (UP_FILL if delta else None)
        for c in ws[ws.max_row]:
            c.font = Font(name="Arial", size=10, bold=r["ps_id"] in watch)
            if fill:
                c.fill = fill
        for col in (8, 9):
            ws.cell(ws.max_row, col).number_format = '+0;-0;"-"'
    ws.auto_filter.ref = f"A1:K{ws.max_row}"


def write_history(wb, entries):
    if "History" not in wb.sheetnames:
        ws = wb.create_sheet("History")
        style_header(ws, ["Timestamp", "PS ID", "Title", "Previous", "Now", "Change"],
                     [18, 11, 60, 10, 10, 10])
    ws = wb["History"]
    for e in entries:
        ws.append([e["time"], e["ps_id"], e["title"], e["prev"], e["now"], e["now"] - e["prev"]])
        for c in ws[ws.max_row]:
            c.font = FONT
    ws.auto_filter.ref = f"A1:F{ws.max_row}"


def write_summary(wb, n_rows, stamp):
    """Summary uses live formulas pointing at the Live sheet, so it stays correct
    even if you sort or filter Live yourself."""
    if "Summary" in wb.sheetnames:
        del wb["Summary"]
    ws = wb.create_sheet("Summary")
    last = n_rows + 1
    rng = f"Live!$G$2:$G${last}"
    items = [
        ("Last updated", stamp),
        ("Source", URL),
        ("Total problem statements", f"=COUNTA(Live!$B$2:$B${last})"),
        ("Total ideas submitted", f"=SUM({rng})"),
        ("Software ideas", f'=SUMIFS({rng},Live!$E$2:$E${last},"Software")'),
        ("Hardware ideas", f'=SUMIFS({rng},Live!$E$2:$E${last},"Hardware")'),
        ("New ideas since last run", f"=SUM(Live!$H$2:$H${last})"),
        ("New ideas today", f"=SUM(Live!$I$2:$I${last})"),
        ("Average ideas per PS", f"=IFERROR(AVERAGE({rng}),0)"),
        ("PS with zero ideas", f"=COUNTIF({rng},0)"),
    ]
    for label, val in items:
        ws.append([label, val])
    for row in ws.iter_rows():
        row[0].font = Font(name="Arial", size=10, bold=True)
        row[1].font = FONT
    ws["B9"].number_format = "0.0"
    ws.column_dimensions["A"].width = 28
    ws.column_dimensions["B"].width = 40

    # Top 10 most and least crowded, via INDEX/MATCH on LARGE/SMALL
    ws["A12"], ws["D12"] = "Most submissions (top 10)", "Fewest submissions (bottom 10)"
    for c in ("A12", "D12"):
        ws[c].font = Font(name="Arial", size=11, bold=True)
    for col, txt in zip("ABDE", ["PS ID", "Ideas", "PS ID", "Ideas"]):
        ws[f"{col}13"] = txt
        ws[f"{col}13"].font, ws[f"{col}13"].fill = HEAD_FONT, HEAD_FILL
    ws.column_dimensions["D"].width = 14
    ws.column_dimensions["E"].width = 10
    for k in range(1, 11):
        r = 13 + k
        ws[f"B{r}"] = f"=LARGE({rng},{k})"
        ws[f"E{r}"] = f"=SMALL({rng},{k})"
        # COUNTIF offset handles ties so the same PS isn't listed twice
        ws[f"A{r}"] = (f"=INDEX(Live!$B$2:$B${last},MATCH(1,INDEX(({rng}=B{r})*"
                       f"(COUNTIF(A$13:A{r - 1},Live!$B$2:$B${last})=0),0),0))")
        ws[f"D{r}"] = (f"=INDEX(Live!$B$2:$B${last},MATCH(1,INDEX(({rng}=E{r})*"
                       f"(COUNTIF(D$13:D{r - 1},Live!$B$2:$B${last})=0),0),0))")
        for c in ("A", "B", "D", "E"):
            ws[f"{c}{r}"].font = FONT


def save_workbook(rows, state, new_history, watch, stamp):
    wb = load_workbook(EXCEL_FILE) if EXCEL_FILE.exists() else Workbook()
    if "Sheet" in wb.sheetnames and len(wb.sheetnames) == 1:
        del wb["Sheet"]
    write_live(wb, rows, state["last"], state["day_start"], watch, stamp)
    write_history(wb, state["pending_history"] + new_history)
    write_summary(wb, len(rows), stamp)
    wb._sheets = [wb["Live"], wb["Summary"], wb["History"]]
    wb.active = 0
    try:
        wb.save(EXCEL_FILE)
        return True
    except PermissionError:
        alt = EXCEL_FILE.with_name(f"{EXCEL_FILE.stem}_{datetime.now():%H%M%S}.xlsx")
        wb.save(alt)
        log(f"{EXCEL_FILE} is open/locked - saved a copy to {alt}; will retry main file next run")
        return False


# --------------------------------------------------------------------------- main
def run_once(watch, cache=None):
    stamp = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
    html = Path(cache).read_text(encoding="utf-8", errors="replace") if cache else fetch_with_retry()
    if html is None:
        log("Could not reach sih.gov.in - skipping this run")
        return False
    try:
        rows = parse(html)
    except RuntimeError as e:
        log(f"Parse error: {e}")
        return False

    state = load_state()
    today = datetime.now().strftime("%Y-%m-%d")
    if state["day"] != today:  # new day -> reset the "today" baseline
        state["day"] = today
        state["day_start"] = state["last"] or {r["ps_id"]: r["ideas"] for r in rows}

    new_history = []
    for r in rows:
        before = state["last"].get(r["ps_id"])
        if before is not None and before != r["ideas"]:
            new_history.append({"time": stamp, "ps_id": r["ps_id"], "title": r["title"],
                                "prev": before, "now": r["ideas"]})

    ok = save_workbook(rows, state, new_history, watch, stamp)
    # If the main file was locked, keep history entries so they land there next time
    state["pending_history"] = [] if ok else state["pending_history"] + new_history
    state["last"] = {r["ps_id"]: r["ideas"] for r in rows}
    save_state(state)

    total = sum(r["ideas"] for r in rows)
    log(f"Updated {len(rows)} PS | total ideas: {total} | {len(new_history)} PS changed")
    for ps in watch:
        hit = next((r for r in rows if r["ps_id"] == ps), None)
        if hit:
            log(f"  watch {ps}: {hit['ideas']} ideas")
    return True


def main():
    global EXCEL_FILE
    ap = argparse.ArgumentParser(description="Track SIH 2026 idea submissions in Excel")
    ap.add_argument("--interval", type=float, default=5, help="minutes between updates (default 5)")
    ap.add_argument("--once", action="store_true", help="run one update and exit")
    ap.add_argument("--watch", nargs="*", default=[], help="PS IDs to highlight, e.g. SIH26011")
    ap.add_argument("--output", default=str(EXCEL_FILE), help="Excel file path")
    ap.add_argument("--cache", help="parse a saved HTML file instead of fetching")
    args = ap.parse_args()

    EXCEL_FILE = Path(args.output)
    watch = {w.upper() for w in args.watch}

    if args.once:
        # exit code 1 on failure so CI (GitHub Actions) shows a clear red step
        sys.exit(0 if run_once(watch, args.cache) else 1)
    log(f"Tracking {URL} every {args.interval:g} min -> {EXCEL_FILE.resolve()}  (Ctrl+C to stop)")
    while True:
        start = time.time()
        try:
            run_once(watch, args.cache)
        except Exception as e:  # never let one bad run kill the loop
            log(f"Unexpected error: {e}")
        time.sleep(max(0, args.interval * 60 - (time.time() - start)))


if __name__ == "__main__":
    try:
        main()
    except KeyboardInterrupt:
        log("Stopped.")
